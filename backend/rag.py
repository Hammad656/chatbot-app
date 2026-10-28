import os
import re
import zipfile
from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, LargeZipFile
from uuid import UUID

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pinecone import Pinecone
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_EXTRACTED_CHARS = 100_000
MAX_PDF_PAGES = 80
MAX_DOCX_UNCOMPRESSED_BYTES = 24 * 1024 * 1024
CHUNK_SIZE = 1_200
CHUNK_OVERLAP = 180
PINECONE_BATCH_SIZE = 100


@dataclass(frozen=True)
class CVChunk:
    text: str
    section: str


@dataclass(frozen=True)
class RetrievedChunk:
    text: str
    section: str
    chunk_index: int
    score: float


@lru_cache(maxsize=1)
def _pinecone_index():
    api_key = os.getenv("PINECONE_API_KEY")
    index_name = os.getenv("PINECONE_INDEX_NAME")
    if not _pinecone_credentials_configured(api_key, index_name):
        raise RuntimeError(
            "Set a real PINECONE_API_KEY and PINECONE_INDEX_NAME in backend/.env."
        )
    return Pinecone(api_key=api_key).Index(index_name)


def _pinecone_credentials_configured(
    api_key: str | None, index_name: str | None
) -> bool:
    placeholders = {
        "your-key",
        "your-pinecone-api-key",
        "replace-with-your-pinecone-api-key",
    }
    return bool(api_key and api_key.strip().lower() not in placeholders and index_name)


def pinecone_is_configured() -> bool:
    return _pinecone_credentials_configured(
        os.getenv("PINECONE_API_KEY"), os.getenv("PINECONE_INDEX_NAME")
    )


def require_pinecone_configuration() -> None:
    if not pinecone_is_configured():
        raise RuntimeError(
            "Set a real PINECONE_API_KEY and PINECONE_INDEX_NAME in backend/.env."
        )


def safe_filename(filename: str) -> str:
    basename = PurePosixPath(filename.replace("\\", "/")).name
    cleaned = "".join(
        character for character in basename if character.isprintable()
    ).strip()
    return cleaned[:255] or "cv"


def _extract_pdf(content: bytes) -> list[tuple[str, str]]:
    reader = PdfReader(BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise ValueError("This PDF is password-protected. Upload an unlocked PDF.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"PDFs must contain {MAX_PDF_PAGES} pages or fewer.")

    sections = []
    extracted_size = 0
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            extracted_size += len(text)
            if extracted_size > MAX_EXTRACTED_CHARS:
                raise ValueError(
                    f"CV text exceeds the {MAX_EXTRACTED_CHARS:,}-character processing limit."
                )
            sections.append((f"Page {page_number}", text))
    return sections


def _extract_docx(content: bytes) -> list[tuple[str, str]]:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if (
                sum(item.file_size for item in archive.infolist())
                > MAX_DOCX_UNCOMPRESSED_BYTES
            ):
                raise ValueError("The DOCX file expands beyond the allowed size.")
    except (BadZipFile, LargeZipFile) as error:
        raise ValueError("The uploaded DOCX file is invalid or unreadable.") from error

    try:
        document = Document(BytesIO(content))
    except (
        BadZipFile,
        KeyError,
        LargeZipFile,
        PackageNotFoundError,
        ValueError,
        XMLSyntaxError,
    ) as error:
        raise ValueError("The uploaded DOCX file is invalid or unreadable.") from error
    current_section = "CV"
    sections: list[tuple[str, str]] = []
    extracted_size = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        extracted_size += len(text)
        if extracted_size > MAX_EXTRACTED_CHARS:
            raise ValueError(
                f"CV text exceeds the {MAX_EXTRACTED_CHARS:,}-character processing limit."
            )
        style_name = paragraph.style.name.lower() if paragraph.style else ""
        if style_name.startswith("heading"):
            current_section = text[:160]
        else:
            sections.append((current_section, text))

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text = " | ".join(cells)
                extracted_size += len(text)
                if extracted_size > MAX_EXTRACTED_CHARS:
                    raise ValueError(
                        f"CV text exceeds the {MAX_EXTRACTED_CHARS:,}-character processing limit."
                    )
                sections.append((current_section, text))
    return sections


def extract_cv_chunks(filename: str, content: bytes) -> list[CVChunk]:
    extension = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("The uploaded file is not a valid PDF.")
        try:
            sections = _extract_pdf(content)
        except (PdfReadError, EOFError, KeyError) as error:
            raise ValueError("The uploaded PDF file is invalid or unreadable.") from error
    elif extension == ".docx":
        if not content.startswith(b"PK"):
            raise ValueError("The uploaded file is not a valid DOCX document.")
        sections = _extract_docx(content)
    else:
        raise ValueError("Upload a PDF or DOCX file.")

    extracted_size = sum(len(text) for _, text in sections)
    if extracted_size > MAX_EXTRACTED_CHARS:
        raise ValueError(
            f"CV text exceeds the {MAX_EXTRACTED_CHARS:,}-character processing limit."
        )
    if not sections:
        raise ValueError(
            "No selectable text was found. Scanned PDFs need OCR before upload."
        )

    chunks: list[CVChunk] = []
    for section, text in sections:
        normalized = re.sub(r"[ \t]+", " ", text)
        normalized = re.sub(r"\s*\n\s*", "\n", normalized).strip()
        start = 0
        while start < len(normalized):
            end = min(start + CHUNK_SIZE, len(normalized))
            if end < len(normalized):
                boundary = max(
                    normalized.rfind(" ", start + int(CHUNK_SIZE * 0.7), end),
                    normalized.rfind("\n", start + int(CHUNK_SIZE * 0.7), end),
                )
                if boundary > start:
                    end = boundary
            text_chunk = normalized[start:end].strip()
            if text_chunk:
                chunks.append(CVChunk(text=text_chunk, section=section[:160]))
            if end >= len(normalized):
                break
            start = max(start + 1, end - CHUNK_OVERLAP)

    return chunks


def _namespace(user_id: UUID) -> str:
    return f"user-{user_id.hex}"


def upsert_cv(user_id: UUID, document_id: str, chunks: list[CVChunk]) -> None:
    index = _pinecone_index()
    namespace = _namespace(user_id)
    for offset in range(0, len(chunks), PINECONE_BATCH_SIZE):
        batch = chunks[offset : offset + PINECONE_BATCH_SIZE]
        documents = [
            {
                "_id": f"{document_id}-{offset + position}",
                "text": chunk.text,
                "document_id": document_id,
                "chunk_index": offset + position,
                "section": chunk.section,
            }
            for position, chunk in enumerate(batch)
        ]
        index.documents.upsert(namespace=namespace, documents=documents)


def retrieve_cv_chunks(
    user_id: UUID, document_id: str, query: str, top_k: int = 10
) -> list[RetrievedChunk]:
    response = _pinecone_index().documents.search(
        namespace=_namespace(user_id),
        top_k=top_k,
        score_by=[{"type": "embed", "fields": ["text"], "query": query}],
        filter={"document_id": {"$eq": document_id}},
        include_fields=["text", "section", "chunk_index"],
    )
    hits = response.matches
    return [
        RetrievedChunk(
            text=str(hit.get("text", "")),
            section=str(hit.get("section", "CV")),
            chunk_index=int(hit.get("chunk_index", index)),
            score=float(hit.score or 0),
        )
        for index, hit in enumerate(hits)
        if hit.get("text")
    ]


def delete_cv(user_id: UUID, document_id: str | None = None) -> None:
    index = _pinecone_index()
    namespace = _namespace(user_id)
    if document_id is None:
        index.documents.delete(namespace=namespace, delete_all=True)
    else:
        index.documents.delete(
            namespace=namespace,
            filter={"document_id": {"$eq": document_id}},
        )
