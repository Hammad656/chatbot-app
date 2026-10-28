import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch
from uuid import UUID

from docx import Document
from pinecone import Document as PineconeDocument

from rag import CVChunk, CHUNK_SIZE, extract_cv_chunks, retrieve_cv_chunks, upsert_cv


class CVChunkingTests(unittest.TestCase):
    def test_extracts_docx_headings_and_chunks_long_sections(self):
        document = Document()
        document.add_heading("Experience", level=1)
        document.add_paragraph("Built software for customers. " * 100)
        stream = BytesIO()
        document.save(stream)

        chunks = extract_cv_chunks("resume.docx", stream.getvalue())

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.section == "Experience" for chunk in chunks))
        self.assertTrue(all(len(chunk.text) <= CHUNK_SIZE for chunk in chunks))
        self.assertIn(chunks[0].text[-100:], chunks[1].text)

    def test_rejects_unsupported_and_invalid_files(self):
        with self.assertRaisesRegex(ValueError, "PDF or DOCX"):
            extract_cv_chunks("resume.txt", b"plain text")
        with self.assertRaisesRegex(ValueError, "valid PDF"):
            extract_cv_chunks("resume.pdf", b"not a pdf")
        with self.assertRaisesRegex(ValueError, "valid DOCX"):
            extract_cv_chunks("resume.docx", b"not a docx")

    def test_reads_pinecone_document_matches_as_retrieved_chunks(self):
        matching_document = PineconeDocument(
            {
                "_id": "document-1",
                "text": "Built distributed services.",
                "section": "Experience",
                "chunk_index": 2,
                "_score": 0.82,
            }
        )
        lower_scoring_document = PineconeDocument(
            {
                "_id": "document-1-3",
                "text": "Managed customer projects and deliverables.",
                "section": "Projects",
                "chunk_index": 3,
                "_score": 0.18,
            }
        )
        search = Mock(
            return_value=SimpleNamespace(
                matches=[matching_document, lower_scoring_document]
            )
        )
        fake_index = SimpleNamespace(documents=SimpleNamespace(search=search))

        with patch("rag._pinecone_index", return_value=fake_index):
            chunks = retrieve_cv_chunks(
                user_id=UUID("00000000-0000-0000-0000-000000000001"),
                document_id="document-1",
                query="experience",
            )

        self.assertEqual(chunks[0].section, "Experience")
        self.assertEqual(chunks[0].chunk_index, 2)
        self.assertEqual(chunks[0].text, "Built distributed services.")
        self.assertEqual(chunks[0].score, 0.82)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[1].text, "Managed customer projects and deliverables.")
        search.assert_called_once()
        self.assertEqual(
            search.call_args.kwargs["filter"],
            {"document_id": {"$eq": "document-1"}},
        )
        self.assertEqual(search.call_args.kwargs["top_k"], 10)
        self.assertEqual(chunks[1].score, 0.18)

    def test_upserts_text_field_and_per_user_namespace(self):
        upsert = Mock()
        fake_index = SimpleNamespace(documents=SimpleNamespace(upsert=upsert))
        user_id = UUID("00000000-0000-0000-0000-000000000001")

        with patch("rag._pinecone_index", return_value=fake_index):
            upsert_cv(
                user_id,
                "document-1",
                [CVChunk(text="Built distributed services.", section="Experience")],
            )

        upsert.assert_called_once_with(
            namespace=f"user-{user_id.hex}",
            documents=[
                {
                    "_id": "document-1-0",
                    "text": "Built distributed services.",
                    "document_id": "document-1",
                    "chunk_index": 0,
                    "section": "Experience",
                }
            ],
        )


if __name__ == "__main__":
    unittest.main()
